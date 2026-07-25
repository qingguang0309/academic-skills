#!/usr/bin/env python3
"""为 .docx 应用 qingguang 的个人风格：
- 页眉：左侧北大 logo（assets/pku_logo.png，0.95in × 0.22in），
  右侧文档大标题（Microsoft YaHei，北大红，右对齐 tab）
- 页眉下方：北大红 0.5pt 分隔线
- 页脚：居中 "第X页/共X页"，用 PAGE / NUMPAGES 域（自动更新）
- 全文字体：正文中文宋体 + 西文 Times New Roman；标题中文微软雅黑 + 西文 Arial
- 全文各级标题统一北大红（#94070A），不引入其他颜色

用法：python apply_style.py <doc.docx> --title "文档大标题" [--logo path] [-o out.docx]
依赖：pip install python-docx --break-system-packages
"""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement

DEFAULT_LOGO = Path(__file__).resolve().parent.parent / "assets" / "pku_logo.png"
LOGO_W_IN, LOGO_H_IN = 0.95, 0.22          # 页眉 logo 尺寸
PKU_RED = "94070A"                          # 北大红（全文唯一强调色）
PKU_RED_RGB = RGBColor(0x94, 0x07, 0x0A)    # 同上，供字体颜色使用
TITLE_FONT = "Microsoft YaHei"
TITLE_SIZE_PT = 9
BODY_LATIN, BODY_EAST = "Times New Roman", "SimSun"   # 正文：西文/中文
HEAD_LATIN, HEAD_EAST = "Arial", "Microsoft YaHei"    # 各级标题：西文/中文
# 覆盖所有级别：总标题 + 副标题 + Heading 1–9（Word 支持到 9 级），一律北大红、不逐级换色
HEADING_STYLES = ("Title", "Subtitle") + tuple(f"Heading {i}" for i in range(1, 10))


def _is_heading_style(name):
    """判断段落样式是否属于标题族（含中文本地化名 “标题 1” 及自定义 Heading*）。"""
    if not name:
        return False
    if name in ("Title", "Subtitle", "标题", "副标题"):
        return True
    return name.startswith("Heading") or name.startswith("标题 ")


def _clear(container):
    for p in list(container.paragraphs):
        p._element.getparent().remove(p._element)
    return container.add_paragraph()


def _field(par, instr):
    """在段落中插入 Word 域（如 PAGE / NUMPAGES）。"""
    run = par.add_run()
    for el, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, f" {instr} "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        e = OxmlElement(el)
        for k, v in attrs.items():
            e.set(qn(k), v)
        if text:
            e.text = text
        run._r.append(e)


def _set_run_fonts(run_or_style, latin, east):
    """同时设置西文字体与中文字体（eastAsia 必须走 XML）。"""
    font = run_or_style.font
    font.name = latin
    rpr = font.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _bottom_border(par, color=PKU_RED, sz_eighth_pt=4):
    """段落下边框：0.5pt = sz 4（单位为 1/8 pt）。"""
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz_eighth_pt))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def _iter_all_paragraphs(doc):
    """遍历正文段落 + 所有表格单元格内的段落（标题偶尔落在表格里）。"""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _redden_heading_runs(par):
    """把标题段落每个 run 直接染北大红，覆盖直排式的行内颜色，
    确保样式级颜色被局部格式盖过时仍统一为北大红。"""
    for run in par.runs:
        run.font.color.rgb = PKU_RED_RGB


def _apply_document_fonts(doc):
    """正文与标题的默认字体+颜色。
    正文（含参考文献等英文内容）随 Normal 落到 Times New Roman，用默认黑色。
    各级标题（Title/Subtitle/Heading 1–9，即所有级别）统一北大红，
    全文除黑色正文外只保留这一种强调色。"""
    _set_run_fonts(doc.styles["Normal"], BODY_LATIN, BODY_EAST)
    # 1) 样式级：字体 + 颜色，覆盖所有存在的标题样式
    for name in HEADING_STYLES:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _set_run_fonts(style, HEAD_LATIN, HEAD_EAST)
        style.font.color.rgb = PKU_RED_RGB
    # 2) 段落级兜底：任何标题族段落，逐 run 直接染红，防止局部格式盖过样式色
    for par in _iter_all_paragraphs(doc):
        if _is_heading_style(par.style.name):
            _redden_heading_runs(par)


def apply_style(doc_path, title, logo_path=DEFAULT_LOGO, out_path=None):
    doc = Document(doc_path)
    _apply_document_fonts(doc)

    for section in doc.sections:
        section.different_first_page_header_footer = False

        # --- 页眉：logo 左 + 标题右（右对齐 tab stop 到右边距处）---
        header_par = _clear(section.header)
        header_par.paragraph_format.tab_stops.clear_all()
        usable = section.page_width - section.left_margin - section.right_margin
        header_par.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)

        logo_run = header_par.add_run()
        logo_run.add_picture(str(logo_path), width=Inches(LOGO_W_IN), height=Inches(LOGO_H_IN))
        header_par.add_run("\t")
        title_run = header_par.add_run(str(title))
        _set_run_fonts(title_run, TITLE_FONT, TITLE_FONT)
        title_run.font.size = Pt(TITLE_SIZE_PT)
        title_run.font.color.rgb = PKU_RED_RGB

        # 页眉下方：北大红 0.5pt 分隔线
        _bottom_border(header_par)

        # --- 页脚：第X页/共X页，居中，用域 ---
        footer_par = _clear(section.footer)
        footer_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_par.add_run("第")
        _field(footer_par, "PAGE")
        footer_par.add_run("页/共")
        _field(footer_par, "NUMPAGES")
        footer_par.add_run("页")

    doc.save(out_path or doc_path)
    print(f"OK: {out_path or doc_path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--title", required=True)
    ap.add_argument("--logo", default=str(DEFAULT_LOGO))
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    apply_style(a.docx, a.title, Path(a.logo), a.out)
